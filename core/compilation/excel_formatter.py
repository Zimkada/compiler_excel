"""
Formatage Excel pour la sortie de compilation
Auteur: GOUNOU N'GOBI Chabi Zimé
Version: 3.2
"""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from typing import List, Optional


# Couleurs Excel
EXCEL_COLORS = {
    "PRIMARY": "217346",  # Vert Excel
    "LIGHT_TEXT": "FFFFFF",  # Blanc
}

# Formats de date Excel
DATE_FORMATS_EXCEL = {
    "FRENCH": "dd/mm/yyyy",
    "AMERICAN": "mm/dd/yyyy",
    "ISO": "yyyy-mm-dd",
    "DATETIME_FRENCH": "dd/mm/yyyy hh:mm:ss",
}


class ExcelFormatter:
    """
    Classe responsable du formatage du fichier Excel de sortie.
    Applique des styles professionnels aux en-têtes et données.
    """

    # Styles prédéfinis
    HEADER_STYLE = {
        'font': Font(bold=True, color=EXCEL_COLORS["LIGHT_TEXT"], size=11),
        'fill': PatternFill(
            start_color=EXCEL_COLORS["PRIMARY"],
            end_color=EXCEL_COLORS["PRIMARY"],
            fill_type='solid'
        ),
        'alignment': Alignment(horizontal='center', vertical='center', wrap_text=True)
    }

    DATA_BORDER = Border(
        left=Side(border_style='thin', color='D0D0D0'),
        right=Side(border_style='thin', color='D0D0D0'),
        top=Side(border_style='thin', color='D0D0D0'),
        bottom=Side(border_style='thin', color='D0D0D0')
    )

    @staticmethod
    def write_headers(worksheet, headers: List[List], start_row: int) -> int:
        """
        Écrit les en-têtes avec style professionnel

        Args:
            worksheet: Feuille de calcul openpyxl
            headers: Liste de listes représentant les en-têtes
            start_row: Ligne de début

        Returns:
            Numéro de la ligne suivante
        """
        current_row = start_row

        for header_row in headers:
            for col_idx, value in enumerate(header_row, 1):
                cell = worksheet.cell(row=current_row, column=col_idx, value=value)
                cell.font = ExcelFormatter.HEADER_STYLE['font']
                cell.fill = ExcelFormatter.HEADER_STYLE['fill']
                cell.alignment = ExcelFormatter.HEADER_STYLE['alignment']
                cell.border = ExcelFormatter.DATA_BORDER

            current_row += 1

        return current_row

    # Au-delà de ce nombre de lignes de données, on n'applique plus de bordure
    # cellule par cellule : sur 50 000 lignes × 12 colonnes = 600 000 cellules,
    # instancier et affecter un objet Border par cellule dominait le temps
    # d'écriture (mesuré : ~7 min). La bordure est purement cosmétique ; au-delà
    # du seuil on l'omet pour rester rapide (le tableau reste lisible, en-têtes
    # stylés et volets figés). En deçà, on conserve le rendu bordé complet.
    BORDER_ROW_LIMIT = 5000

    @staticmethod
    def write_data(worksheet, data: List[List], start_row: int,
                   date_format: str = "FRENCH") -> int:
        """
        Écrit les données rapidement, avec format de date et bordures adaptatives.

        Chemin rapide : les lignes sont ajoutées via ``worksheet.append`` (bien
        plus rapide que ``worksheet.cell(...)`` répété). Les bordures ne sont
        appliquées que sous ``BORDER_ROW_LIMIT`` lignes (au-delà, elles coûtent
        trop cher pour un gain purement cosmétique). Le format de date n'est
        posé que sur les cellules réellement datetime.

        Args:
            worksheet: Feuille de calcul openpyxl
            data: Liste de listes représentant les données
            start_row: Ligne de début
            date_format: Format de date (FRENCH, AMERICAN, ISO)

        Returns:
            Numéro de la ligne suivante
        """
        excel_date_format = DATE_FORMATS_EXCEL.get(date_format, "dd/mm/yyyy")
        apply_borders = len(data) <= ExcelFormatter.BORDER_ROW_LIMIT
        border = ExcelFormatter.DATA_BORDER

        current_row = start_row
        for row_data in data:
            # append place la ligne d'un coup à la fin (rapide). On ne repasse
            # sur les cellules que si un style ponctuel est nécessaire.
            worksheet.append(list(row_data))

            if apply_borders or any(isinstance(v, datetime) for v in row_data):
                for col_idx, value in enumerate(row_data, 1):
                    cell = worksheet.cell(row=current_row, column=col_idx)
                    if apply_borders:
                        cell.border = border
                    if isinstance(value, datetime):
                        cell.number_format = excel_date_format

            current_row += 1

        return current_row

    # Nombre de lignes échantillonnées pour estimer la largeur des colonnes.
    # Parcourir toutes les cellules (worksheet.columns matérialise TOUTE la
    # feuille) coûtait autant que l'écriture elle-même sur les gros fichiers.
    # Un échantillon de tête suffit à dimensionner correctement les colonnes.
    WIDTH_SAMPLE_ROWS = 200

    @staticmethod
    def adjust_column_widths(worksheet, min_width: int = 10, max_width: int = 50):
        """
        Ajuste la largeur des colonnes à partir d'un ÉCHANTILLON de lignes.

        On lit au plus ``WIDTH_SAMPLE_ROWS`` lignes via iter_rows (accès
        séquentiel rapide) au lieu de matérialiser toute la feuille via
        ``worksheet.columns``. La largeur reste représentative : les libellés
        d'en-tête et les premières lignes déterminent l'essentiel.

        Args:
            worksheet: Feuille de calcul openpyxl
            min_width: Largeur minimale
            max_width: Largeur maximale
        """
        max_col = worksheet.max_column or 0
        if max_col == 0:
            return
        max_lengths = [0] * max_col

        for row in worksheet.iter_rows(
            min_row=1, max_row=min(worksheet.max_row, ExcelFormatter.WIDTH_SAMPLE_ROWS)
        ):
            for cell in row:
                if cell.value is not None:
                    col_i = cell.column - 1
                    if 0 <= col_i < max_col:
                        length = len(str(cell.value))
                        if length > max_lengths[col_i]:
                            max_lengths[col_i] = length

        for col_i, max_length in enumerate(max_lengths, start=1):
            adjusted_width = max(min_width, min(max_length + 2, max_width))
            worksheet.column_dimensions[get_column_letter(col_i)].width = adjusted_width

    @staticmethod
    def freeze_header(worksheet, freeze_row: int):
        """
        Fige les volets à la ligne spécifiée

        Args:
            worksheet: Feuille de calcul openpyxl
            freeze_row: Ligne où figer (les lignes au-dessus seront figées)
        """
        if freeze_row > 1:
            worksheet.freeze_panes = worksheet.cell(row=freeze_row, column=1)

    @staticmethod
    def apply_alternating_rows(worksheet, start_row: int, end_row: int,
                               color: str = "F2F2F2"):
        """
        Applique une couleur alternée aux lignes de données

        Args:
            worksheet: Feuille de calcul openpyxl
            start_row: Ligne de début
            end_row: Ligne de fin
            color: Couleur hexadécimale (sans #)
        """
        fill = PatternFill(start_color=color, end_color=color, fill_type='solid')

        for row_idx in range(start_row, end_row + 1):
            if row_idx % 2 == 0:  # Lignes paires
                for cell in worksheet[row_idx]:
                    if cell.fill.fill_type is None:
                        cell.fill = fill
