"""
Tests de l'écriture XLSX rapide (E1).

L'écriture cellule par cellule avec bordure (+ largeurs calculées sur toute la
feuille) faisait ~7 min pour 50 000 lignes. Le chemin rapide écrit via append,
n'applique les bordures que sous un seuil, et échantillonne les largeurs.

On vérifie l'INTÉGRITÉ (aucune donnée perdue, styles clés conservés) et le
comportement adaptatif des bordures. Un benchmark chronométré est marqué `slow`
(exécuté à la demande) pour ne pas alourdir la suite.
"""

import time

import openpyxl
import pandas as pd
import pytest

from core.compilation.excel_formatter import ExcelFormatter


def _write(tmp_path, n_rows, n_cols=12):
    headers = [[f'Col{c}' for c in range(n_cols)]]
    data = [[f'v{i}_{c}' if c % 3 else i * c for c in range(n_cols)]
            for i in range(n_rows)]
    wb = openpyxl.Workbook()
    ws = wb.active
    cur = ExcelFormatter.write_headers(ws, headers, start_row=1)
    ExcelFormatter.write_data(ws, data, start_row=cur)
    ExcelFormatter.adjust_column_widths(ws, min_width=10, max_width=50)
    ExcelFormatter.freeze_header(ws, len(headers) + 1)
    out = tmp_path / 'out.xlsx'
    wb.save(out)
    return ws, out


# --------------------------------------------------------------------------- #
# Intégrité
# --------------------------------------------------------------------------- #
class TestIntegrity:
    def test_no_data_loss_small(self, tmp_path):
        ws, out = _write(tmp_path, 50)
        back = pd.read_excel(out, header=None)
        assert len(back) == 51  # 1 en-tête + 50 données

    def test_no_data_loss_large(self, tmp_path):
        ws, out = _write(tmp_path, 12000)
        back = pd.read_excel(out, header=None)
        assert len(back) == 12001

    def test_header_style_preserved(self, tmp_path):
        ws, _ = _write(tmp_path, 100)
        assert ws['A1'].fill.fill_type == 'solid'
        assert ws['A1'].font.bold is True

    def test_freeze_panes_set(self, tmp_path):
        ws, _ = _write(tmp_path, 100)
        assert ws.freeze_panes == 'A2'

    def test_column_widths_applied(self, tmp_path):
        ws, _ = _write(tmp_path, 100)
        # Toutes les colonnes écrites ont une largeur dans [min, max].
        for c in range(1, 13):
            from openpyxl.utils import get_column_letter
            w = ws.column_dimensions[get_column_letter(c)].width
            assert 10 <= w <= 50

    def test_values_intact_at_boundaries(self, tmp_path):
        """Première et dernière ligne de données correctes (pas de décalage
        introduit par append après les en-têtes)."""
        ws, out = _write(tmp_path, 30)
        back = pd.read_excel(out, header=None).values.tolist()
        # back[0] = en-tête, back[1] = 1re donnée (i=0), back[30] = dernière (i=29)
        assert back[1][0] == 0        # i*c avec c=0 -> 0
        assert back[30][1] == 'v29_1'  # i=29, c=1


# --------------------------------------------------------------------------- #
# Bordures adaptatives
# --------------------------------------------------------------------------- #
class TestAdaptiveBorders:
    def test_borders_present_below_limit(self, tmp_path):
        ws, _ = _write(tmp_path, 100)
        # Sous le seuil : les données portent une bordure.
        assert ws['A2'].border.left.style is not None

    def test_borders_omitted_above_limit(self, tmp_path):
        n = ExcelFormatter.BORDER_ROW_LIMIT + 100
        ws, _ = _write(tmp_path, n)
        # Au-delà du seuil : pas de bordure sur les données (perf), mais
        # l'en-tête reste stylé et les données sont intactes.
        assert ws['A2'].border.left.style is None
        assert ws['A1'].fill.fill_type == 'solid'

    def test_dates_formatted_even_above_limit(self, tmp_path):
        """Le format de date est appliqué même au-delà du seuil de bordures."""
        from datetime import datetime
        headers = [['Nom', 'Date']]
        n = ExcelFormatter.BORDER_ROW_LIMIT + 10
        data = [[f'P{i}', datetime(2025, 1, 1)] for i in range(n)]
        wb = openpyxl.Workbook()
        ws = wb.active
        cur = ExcelFormatter.write_headers(ws, headers, 1)
        ExcelFormatter.write_data(ws, data, cur, date_format='FRENCH')
        assert ws.cell(row=cur, column=2).number_format == 'dd/mm/yyyy'


# --------------------------------------------------------------------------- #
# Benchmark (lent, à la demande)
# --------------------------------------------------------------------------- #
@pytest.mark.slow
def test_write_50k_under_60s(tmp_path):
    """50 000 lignes doivent s'écrire en moins de 60 s (avant : ~450 s)."""
    t0 = time.time()
    _write(tmp_path, 50000)
    dt = time.time() - t0
    assert dt < 60, f"écriture 50k trop lente: {dt:.1f}s"
