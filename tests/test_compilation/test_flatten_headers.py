"""
Tests adverses de l'aplatissement des en-têtes multi-lignes (défaut D3).

Un en-tête réparti sur plusieurs lignes (catégorie + sous-catégorie, ex.
« NOMBRE DE CAS » au-dessus de « 1er cycle » / « 2nd cycle ») doit être fusionné
en un seul libellé par colonne (« NOMBRE DE CAS - 1er cycle »), produisant une
unique ligne d'en-tête propre. Les fusions verticales recopiées ne doivent pas
donner « DEP - DEP » (déduplication).
"""

import openpyxl
import pandas as pd

from core.compilation import ExcelCompiler, CompilationOptions


def _make_xlsx(path, rows, merges=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row, 1):
            if val is not None:
                ws.cell(ri, ci, val)
    for m in (merges or []):
        ws.merge_cells(m)
    wb.save(path)
    return str(path)


def _compile(tmp_path, src, header_rows, flatten=True):
    opts = CompilationOptions(
        auto_detect_structure=False,
        manual_header_start_row=1,
        manual_header_rows=header_rows,
        unmerge_cells=True,
        flatten_multiindex_headers=flatten,
        remove_empty_rows=True,
    )
    out = str(tmp_path / "out.xlsx")
    ExcelCompiler(opts).compile_files([src], out)
    return pd.read_excel(out, header=None)


class TestFlattenUnit:
    def test_flatten_dedups_repeated_vertical_value(self):
        # Fusion verticale recopiée : « DEP » sur deux lignes -> « DEP », pas « DEP - DEP ».
        comp = ExcelCompiler(CompilationOptions())
        flat = comp._flatten_headers([["DEP", "CAS"], ["DEP", "1er cycle"]])
        assert flat == ["DEP", "CAS - 1er cycle"]

    def test_flatten_ignores_empty_and_nan(self):
        comp = ExcelCompiler(CompilationOptions())
        flat = comp._flatten_headers([["A", None, ""], ["", "B", float("nan")]])
        assert flat == ["A", "B", ""]

    def test_flatten_single_row_is_passthrough(self):
        comp = ExcelCompiler(CompilationOptions())
        flat = comp._flatten_headers([["Nom", "Age"]])
        assert flat == ["Nom", "Age"]


class TestFlattenEndToEnd:
    def test_two_line_header_becomes_one_clean_row(self, tmp_path):
        f = _make_xlsx(tmp_path / "h2.xlsx", [
            ["Region", "NOMBRE DE CAS", None, "AUTEURS", None],
            ["Region", "1er cycle", "2nd cycle", "Ens", "Elv"],
            ["Nord", 5, 3, 2, 1],
        ], merges=["A1:A2", "B1:C1", "D1:E1"])
        df = _compile(tmp_path, f, header_rows=2, flatten=True)
        # Une seule ligne d'en-tête, libellés concaténés et Region dédupliqué.
        assert list(df.iloc[0]) == [
            "Region", "NOMBRE DE CAS - 1er cycle", "NOMBRE DE CAS - 2nd cycle",
            "AUTEURS - Ens", "AUTEURS - Elv",
        ]
        # Les données suivent immédiatement (pas de 2e ligne d'en-tête).
        assert list(df.iloc[1]) == ["Nord", 5, 3, 2, 1]

    def test_disabled_keeps_raw_header_rows(self, tmp_path):
        f = _make_xlsx(tmp_path / "h2.xlsx", [
            ["Region", "NOMBRE DE CAS", None],
            ["Region", "1er cycle", "2nd cycle"],
            ["Nord", 5, 3],
        ], merges=["A1:A2", "B1:C1"])
        df = _compile(tmp_path, f, header_rows=2, flatten=False)
        # Comportement d'origine : deux lignes d'en-tête brutes conservées.
        assert df.iloc[0, 1] == "NOMBRE DE CAS"
        assert df.iloc[1, 1] == "1er cycle"
        assert list(df.iloc[2]) == ["Nord", 5, 3]

    def test_single_line_header_unaffected(self, tmp_path):
        f = _make_xlsx(tmp_path / "h1.xlsx", [
            ["Nom", "Age"],
            ["Alice", 30],
        ])
        df = _compile(tmp_path, f, header_rows=1, flatten=True)
        assert list(df.iloc[0]) == ["Nom", "Age"]
        assert list(df.iloc[1]) == ["Alice", 30]
