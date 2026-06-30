"""
Tests adverses des lignes de sous-total / total (défaut D4).

La détection est déterministe (mots-clés en MOTS ENTIERS, sans accents/casse).
Par défaut les lignes de total sont exclues (anti double comptage) ; sinon elles
sont conservées et, si demandé, marquées dans une colonne « Type de ligne ».
On vérifie surtout l'absence de FAUX POSITIFS (un établissement « CEG TOTALITE »
ne doit jamais être pris pour un total).
"""

import openpyxl
import pandas as pd

from core.compilation import ExcelCompiler, CompilationOptions
from core.compilation.subtotal_detector import (
    classify_row, is_total_row,
    ROW_KIND_DETAIL, ROW_KIND_SUBTOTAL, ROW_KIND_TOTAL,
)


def _make_xlsx(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row, 1):
            if val is not None:
                ws.cell(ri, ci, val)
    wb.save(path)
    return str(path)


class TestClassifyUnit:
    def test_total_beats_subtotal(self):
        # « ENSEMBLE DEPARTEMENT » contient « ENSEMBLE » mais doit primer en « total ».
        assert classify_row(["ENSEMBLE DEPARTEMENT (ALIBORI)", None, 1]) == ROW_KIND_TOTAL

    def test_subtotal_keyword(self):
        assert classify_row(["x", "ENSEMBLE COMMUNE", "ENSEMBLE COMMUNE", 4637]) == ROW_KIND_SUBTOTAL

    def test_total_general(self):
        assert classify_row(["total general", None, 9]) == ROW_KIND_TOTAL

    def test_accents_and_case_insensitive(self):
        assert classify_row(["Énsemble Commune"]) == ROW_KIND_SUBTOTAL

    def test_no_false_positive_substring(self):
        # Pièges : mots contenant un mot-clé en sous-chaîne -> doivent rester détail.
        assert classify_row(["ALIBORI", "BANIKOARA", "CEG TOTALITE", 10]) == ROW_KIND_DETAIL
        assert classify_row(["X", "Y", "CAPITAL HUMAIN", 5]) == ROW_KIND_DETAIL
        assert classify_row(["RASSEMBLEMENT SCOLAIRE", 3]) == ROW_KIND_DETAIL

    def test_numeric_cells_ignored(self):
        assert classify_row([1, 2, 3, 4]) == ROW_KIND_DETAIL

    def test_plural_totaux_detected(self):
        # « TOTAUX » (pluriel) doit aussi être reconnu (régression audit S2).
        assert classify_row(["TOTAUX", 99]) == ROW_KIND_SUBTOTAL

    def test_punctuation_glued_keyword(self):
        # Ponctuation collée au mot-clé : reste détecté (\\b borde le mot).
        assert classify_row(["Total :", 5]) == ROW_KIND_SUBTOTAL
        assert classify_row(["S/TOTAL", 5]) == ROW_KIND_SUBTOTAL

    def test_keyword_glued_to_digits_is_detail(self):
        # « TOTAL2024 » n'est pas le mot « TOTAL » -> pas exclu (régression audit S3).
        assert classify_row(["TOTAL2024", 5]) == ROW_KIND_DETAIL

    def test_is_total_row_helper(self):
        assert is_total_row(["ENSEMBLE COMMUNE", 1]) is True
        assert is_total_row(["CEG NORMAL", 1]) is False


class TestEndToEnd:
    ROWS = [
        ["Region", "Etab", "Cas"],
        ["Nord", "CEG A", 10],
        ["Nord", "CEG TOTALITE", 20],          # piège : ne doit PAS être exclu
        ["Nord", "ENSEMBLE COMMUNE", 30],       # sous-total
        ["TOTAL GENERAL", None, 60],            # total
    ]

    def _compile(self, tmp_path, **opt):
        f = _make_xlsx(tmp_path / "s.xlsx", self.ROWS)
        opts = CompilationOptions(
            auto_detect_structure=False,
            manual_header_start_row=1, manual_header_rows=1,
            unmerge_cells=True, remove_empty_rows=True, **opt,
        )
        out = str(tmp_path / "out.xlsx")
        res = ExcelCompiler(opts).compile_files([f], out)
        return pd.read_excel(out, header=None), res

    def test_default_excludes_totals_keeps_trap(self, tmp_path):
        df, res = self._compile(tmp_path, drop_subtotal_rows=True)
        etabs = [str(v) for v in df.iloc[1:, 1]]
        # Les 2 lignes de total sont parties, le piège « CEG TOTALITE » reste.
        assert "CEG TOTALITE" in etabs
        assert all("ENSEMBLE" not in e.upper() for e in etabs)
        assert res.file_results[0].subtotal_rows == 2
        assert any("2 ligne(s) de total exclues" in w for w in res.warnings)

    def test_keep_and_mark(self, tmp_path):
        df, res = self._compile(tmp_path, drop_subtotal_rows=False, mark_subtotal_rows=True)
        assert df.iloc[0, 0] == "Type de ligne"
        kinds = [str(v) for v in df.iloc[1:, 0]]
        assert kinds.count("détail") == 2      # CEG A + CEG TOTALITE
        assert kinds.count("sous-total") == 1
        assert kinds.count("total") == 1

    def test_keep_without_mark(self, tmp_path):
        df, res = self._compile(tmp_path, drop_subtotal_rows=False, mark_subtotal_rows=False)
        # Pas de colonne marqueur : en-tête inchangé, totaux présents.
        assert df.iloc[0, 0] == "Region"
        assert len(df) == 1 + 4  # en-tête + 4 lignes de données conservées

    def test_preview_matches_compilation_with_exclusion(self, tmp_path):
        # Régression audit S4 : l'aperçu doit compter les lignes APRÈS exclusion
        # des sous-totaux, sinon il annonce plus de lignes que la sortie réelle.
        f = _make_xlsx(tmp_path / "s.xlsx", self.ROWS)
        opts = CompilationOptions(
            auto_detect_structure=False, manual_header_start_row=1,
            manual_header_rows=1, unmerge_cells=True, drop_subtotal_rows=True,
        )
        comp = ExcelCompiler(opts)
        prev = comp.preview_detection([f])[0]
        out = str(tmp_path / "out.xlsx")
        res = comp.compile_files([f], out)
        assert prev.data_row_count == res.file_results[0].rows_added
