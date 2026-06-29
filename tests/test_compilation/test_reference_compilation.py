"""
Tests de non-régression du moteur de compilation en mode référence.

Ce mode (use_reference_mode=True) est celui réellement activé par défaut
dans l'interface. Ces tests verrouillent l'exactitude des DONNÉES produites,
en particulier les bugs corrigés:

- Perte silencieuse des lignes situées après une ligne vide isolée
  (le tableau était tronqué à la première ligne vide).
- Tri non robuste aux types mixtes (TypeError avalé) et colonne hors limites.
- repeat_headers mélangeant des en-têtes dans des données triées/dédupliquées.
- Format de date "date et heure" non câblé.
"""

import openpyxl
import pandas as pd
import pytest

from core.compilation import ExcelCompiler, CompilationOptions, OutputFormat
from core.compilation.compilation_models import FilenameOption, DateFormat


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def _make_xlsx(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)


def _compile(tmp_path, sources, **option_overrides):
    """Compile des listes de lignes en mode référence et renvoie (result, rows)."""
    paths = []
    for i, rows in enumerate(sources):
        p = tmp_path / f"src_{i}.xlsx"
        _make_xlsx(p, rows)
        paths.append(str(p))

    option_overrides.setdefault("remove_empty_rows", True)
    options = CompilationOptions(
        use_reference_mode=True,
        auto_detect_structure=False,
        reference_header_row=1,
        reference_header_lines=1,
        filename_option=FilenameOption.NONE,
        **option_overrides,
    )
    out = tmp_path / "out.xlsx"
    result = ExcelCompiler(options).compile_files(paths, str(out), OutputFormat.XLSX)
    rows = pd.read_excel(out, header=None).values.tolist()
    return result, rows


def _data_names(rows):
    """Noms (1re colonne) des lignes de données, hors en-tête et lignes vides."""
    names = []
    for r in rows:
        v = r[0]
        if v in (None, "Nom", "") or (isinstance(v, float) and pd.isna(v)):
            continue
        names.append(v)
    return names


# --------------------------------------------------------------------------- #
# Fin de tableau / lignes vides (bug de troncature silencieuse)
# --------------------------------------------------------------------------- #
class TestDataEndDetection:
    def test_clean_file_keeps_all_rows(self, tmp_path):
        _, rows = _compile(tmp_path, [[
            ["Nom", "Age"], ["Alice", 30], ["Bob", 25],
            ["Carol", 40], ["Dan", 22], ["Eve", 35],
        ]], remove_empty_rows=False)
        assert _data_names(rows) == ["Alice", "Bob", "Carol", "Dan", "Eve"]

    def test_isolated_empty_row_does_not_truncate(self, tmp_path):
        """Une ligne vide isolée ne doit pas faire perdre les lignes suivantes."""
        _, rows = _compile(tmp_path, [[
            ["Nom", "Age"], ["Alice", 30], ["Bob", 25],
            [None, None], ["Carol", 40], ["Dan", 22],
        ]])
        assert _data_names(rows) == ["Alice", "Bob", "Carol", "Dan"]

    def test_two_empty_rows_below_threshold_keep_data(self, tmp_path):
        """2 lignes vides consécutives (< seuil de 3) ne coupent pas le tableau."""
        _, rows = _compile(tmp_path, [[
            ["Nom", "Age"], ["Alice", 30],
            [None, None], [None, None],
            ["Bob", 25], ["Carol", 40],
        ]])
        assert _data_names(rows) == ["Alice", "Bob", "Carol"]

    def test_block_of_empty_rows_cuts_trailing_content(self, tmp_path):
        """Un vrai bas de tableau (3+ vides) coupe bien le contenu de pied de page."""
        _, rows = _compile(tmp_path, [[
            ["Nom", "Age"], ["Alice", 30], ["Bob", 25],
            [None, None], [None, None], [None, None],
            ["SIGNATURE: M. X", None], ["Date: 2025", None],
        ]])
        assert _data_names(rows) == ["Alice", "Bob"]

    def test_batch_with_gap_in_one_file(self, tmp_path):
        """En batch, un fichier troué ne doit pas perdre ses lignes finales."""
        _, rows = _compile(tmp_path, [
            [["Nom", "Age"], ["A1", 1], ["A2", 2], ["A3", 3]],
            [["Nom", "Age"], ["B1", 1], [None, None], ["B2", 2], ["B3", 3]],
        ])
        assert _data_names(rows) == ["A1", "A2", "A3", "B1", "B2", "B3"]


# --------------------------------------------------------------------------- #
# Tri (robustesse aux types + colonne hors limites)
# --------------------------------------------------------------------------- #
class TestSorting:
    def test_sort_mixed_types_does_not_crash(self, tmp_path):
        """Types mixtes int/str: nombres triés croissants, pas de TypeError avalé."""
        result, rows = _compile(tmp_path, [[
            ["Nom", "Age"], ["Charlie", 30], ["Alice", "inconnu"],
            ["Bob", 40], ["Dan", 5],
        ]], sort_data=True, sort_column=1)
        data = [r for r in rows if r[0] not in (None, "Nom")]
        nums = [r[1] for r in data if isinstance(r[1], (int, float)) and not pd.isna(r[1])]
        assert nums == sorted(nums)

    def test_sort_alpha_case_insensitive(self, tmp_path):
        _, rows = _compile(tmp_path, [[
            ["Nom", "Age"], ["Charlie", 3], ["alice", 2], ["Bob", 4],
        ]], sort_data=True, sort_column=0)
        assert _data_names(rows) == ["alice", "Bob", "Charlie"]

    def test_sort_out_of_bounds_warns_and_preserves_order(self, tmp_path):
        result, rows = _compile(tmp_path, [[
            ["Nom", "Age"], ["Charlie", 3], ["Alice", 2], ["Bob", 4],
        ]], sort_data=True, sort_column=25)
        assert any("hors limites" in w for w in result.warnings)
        assert _data_names(rows) == ["Charlie", "Alice", "Bob"]


# --------------------------------------------------------------------------- #
# repeat_headers et ses interactions
# --------------------------------------------------------------------------- #
class TestRepeatHeaders:
    def test_repeat_headers_alone_repeats_between_files(self, tmp_path):
        _, rows = _compile(tmp_path, [
            [["Nom", "Age"], ["Zoe", 1], ["Amy", 2]],
            [["Nom", "Age"], ["Max", 3], ["Bea", 4]],
        ], repeat_headers=True)
        header_count = sum(1 for r in rows if str(r[0]) == "Nom")
        assert header_count == 2  # en-tête initial + répété avant le 2e fichier

    def test_repeat_headers_disabled_when_sorting(self, tmp_path):
        """repeat_headers + tri => répétition désactivée, aucun en-tête dans les données."""
        result, rows = _compile(tmp_path, [
            [["Nom", "Age"], ["Zoe", 1], ["Amy", 2]],
            [["Nom", "Age"], ["Max", 3], ["Bea", 4]],
        ], repeat_headers=True, sort_data=True, sort_column=0)
        header_in_data = sum(1 for r in rows[1:] if str(r[0]) == "Nom")
        assert header_in_data == 0
        assert any("en-t" in w.lower() for w in result.warnings)
        assert [r[0] for r in rows[1:]] == ["Amy", "Bea", "Max", "Zoe"]

    def test_options_object_not_mutated(self, tmp_path):
        """Le garde-fou ne doit pas muter l'objet options de l'appelant.

        Réutiliser le même CompilationOptions pour une 2e compilation doit
        conserver repeat_headers=True (régression: il était mis à False).
        """
        f1 = tmp_path / "f1.xlsx"
        f2 = tmp_path / "f2.xlsx"
        _make_xlsx(f1, [["Nom", "Age"], ["Zoe", 1], ["Amy", 2]])
        _make_xlsx(f2, [["Nom", "Age"], ["Max", 3], ["Bea", 4]])

        options = CompilationOptions(
            use_reference_mode=True, auto_detect_structure=False,
            reference_header_row=1, reference_header_lines=1,
            filename_option=FilenameOption.NONE, remove_empty_rows=True,
            repeat_headers=True, sort_data=True, sort_column=0,
        )
        out1 = tmp_path / "o1.xlsx"
        ExcelCompiler(options).compile_files([str(f1), str(f2)], str(out1), OutputFormat.XLSX)

        assert options.repeat_headers is True  # non muté

        # 2e compilation sans tri: repeat_headers doit de nouveau s'appliquer
        options.sort_data = False
        out2 = tmp_path / "o2.xlsx"
        ExcelCompiler(options).compile_files([str(f1), str(f2)], str(out2), OutputFormat.XLSX)
        rows = pd.read_excel(out2, header=None).values.tolist()
        header_count = sum(1 for r in rows if str(r[0]) == "Nom")
        assert header_count == 2


# --------------------------------------------------------------------------- #
# Formats de date
# --------------------------------------------------------------------------- #
class TestDateFormats:
    @pytest.mark.parametrize("fmt,expected", [
        (DateFormat.FRENCH, "dd/mm/yyyy"),
        (DateFormat.AMERICAN, "mm/dd/yyyy"),
        (DateFormat.ISO, "yyyy-mm-dd"),
        (DateFormat.DATETIME_FRENCH, "dd/mm/yyyy hh:mm:ss"),
    ])
    def test_date_number_format_applied(self, tmp_path, fmt, expected):
        from datetime import datetime
        src = tmp_path / "dates.xlsx"
        _make_xlsx(src, [["Event", "Quand"],
                         ["A", datetime(2025, 3, 8, 14, 30, 15)]])
        out = tmp_path / "out.xlsx"
        options = CompilationOptions(
            use_reference_mode=True, auto_detect_structure=False,
            reference_header_row=1, reference_header_lines=1,
            filename_option=FilenameOption.NONE, remove_empty_rows=True,
            date_format=fmt,
        )
        ExcelCompiler(options).compile_files([str(src)], str(out), OutputFormat.XLSX)
        cell = openpyxl.load_workbook(out).active.cell(row=2, column=2)
        assert cell.number_format == expected
