"""
Tests adverses de la dé-fusion des cellules (core.compilation.merge_handler).

On fabrique des fichiers piégés couvrant chaque type de fusion observé sur les
vrais fichiers : verticale, horizontale, mixte, imbriquée, formules, et absence
de fusion. La dé-fusion doit propager la valeur du coin haut-gauche sur toute la
plage, sans jamais inventer de donnée.
"""

import openpyxl
import pandas as pd

from core.compilation.merge_handler import load_with_unmerge


def _make_xlsx(path, rows, merges=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row, 1):
            if val is not None:
                ws.cell(ri, ci, val)
    for m in (merges or []):
        ws.merge_cells(m)
    wb.save(path)
    return str(path)


class TestVerticalMerge:
    def test_fill_down_propagates_value(self, tmp_path):
        # "ALIBORI" fusionné verticalement sur A1:A3.
        f = _make_xlsx(tmp_path / "v.xlsx", [
            ["ALIBORI", "CEG A", 10],
            [None, "CEG B", 20],
            [None, "CEG C", 30],
        ], merges=["A1:A3"])
        df = load_with_unmerge(f)
        assert df.iat[0, 0] == "ALIBORI"
        assert df.iat[1, 0] == "ALIBORI"
        assert df.iat[2, 0] == "ALIBORI"
        # Les autres colonnes restent intactes.
        assert list(df.iloc[:, 1]) == ["CEG A", "CEG B", "CEG C"]

    def test_two_vertical_blocks_independent(self, tmp_path):
        f = _make_xlsx(tmp_path / "v2.xlsx", [
            ["ALIBORI", "BANIKOARA"],
            [None, None],
            ["BORGOU", "PARAKOU"],
            [None, None],
        ], merges=["A1:A2", "A3:A4"])
        df = load_with_unmerge(f)
        assert list(df.iloc[:, 0]) == ["ALIBORI", "ALIBORI", "BORGOU", "BORGOU"]


class TestHorizontalMerge:
    def test_fill_across_propagates_value(self, tmp_path):
        # "AUTEURS" fusionné horizontalement sur A1:C1.
        f = _make_xlsx(tmp_path / "h.xlsx", [
            ["AUTEURS", None, None],
            ["Ens", "Elv", "Etu"],
        ], merges=["A1:C1"])
        df = load_with_unmerge(f)
        assert list(df.iloc[0, :3]) == ["AUTEURS", "AUTEURS", "AUTEURS"]
        # La ligne suivante (sous-titres) n'est pas touchée.
        assert list(df.iloc[1, :3]) == ["Ens", "Elv", "Etu"]


class TestMixedAndNested:
    def test_vertical_and_horizontal_together(self, tmp_path):
        f = _make_xlsx(tmp_path / "mix.xlsx", [
            ["DEP", "AUTEURS", None, None],
            [None, "Ens", "Elv", "Etu"],
            [None, 1, 2, 3],
        ], merges=["A1:A2", "B1:D1"])
        df = load_with_unmerge(f)
        assert df.iat[1, 0] == "DEP"           # vertical propagé
        assert list(df.iloc[0, 1:4]) == ["AUTEURS", "AUTEURS", "AUTEURS"]  # horizontal
        # La donnée n'est pas écrasée.
        assert list(df.iloc[2, 1:4]) == [1, 2, 3]

    def test_block_merge_2d(self, tmp_path):
        # Fusion 2D A1:B2 (rare mais possible).
        f = _make_xlsx(tmp_path / "block.xlsx", [
            ["TITRE", None, "X"],
            [None, None, "Y"],
        ], merges=["A1:B2"])
        df = load_with_unmerge(f)
        assert df.iat[0, 0] == "TITRE"
        assert df.iat[0, 1] == "TITRE"
        assert df.iat[1, 0] == "TITRE"
        assert df.iat[1, 1] == "TITRE"
        assert df.iat[0, 2] == "X"
        assert df.iat[1, 2] == "Y"


class TestEdgeCases:
    def test_no_merge_is_noop(self, tmp_path):
        f = _make_xlsx(tmp_path / "plain.xlsx", [
            ["Nom", "Age"],
            ["Alice", 30],
            ["Bob", 25],
        ])
        got = load_with_unmerge(f)
        expected = pd.read_excel(f, header=None)
        # Mêmes dimensions et mêmes valeurs (on ne contraint pas le dtype exact).
        assert got.shape == expected.shape
        assert got.fillna("").values.tolist() == expected.fillna("").values.tolist()

    def test_empty_merge_source_is_noop(self, tmp_path):
        # Fusion dont le coin haut-gauche est vide : rien à propager, pas de crash.
        # (Note : fusionner A1:A2 écrase A2 dans le fichier lui-même — openpyxl
        # ne conserve que le coin haut-gauche. La donnée vide reste vide.)
        f = _make_xlsx(tmp_path / "empty.xlsx", [
            [None, "garde"],
            [None, "moi"],
        ], merges=["A1:A2"])
        df = load_with_unmerge(f)
        # Colonne A reste vide (rien à propager) ; colonne B intacte.
        assert pd.isna(df.iat[0, 0]) or df.iat[0, 0] is None
        assert pd.isna(df.iat[1, 0]) or df.iat[1, 0] is None
        assert list(df.iloc[:, 1]) == ["garde", "moi"]

    def test_uncached_formula_no_regression_vs_pandas(self, tmp_path):
        # Fichier généré sans Excel : la formule n'a pas de valeur cachée.
        # La voie rapide (read_only) se comporte alors comme pd.read_excel
        # (cellule vide), SANS régression : les valeurs littérales restent
        # correctes. Le rattrapage formule-texte n'intervient que dans le
        # chemin de repli openpyxl (cf. test dédié ci-dessous).
        f = str(tmp_path / "formula.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = "=SUM(A1:A2)"
        ws["B3"] = "total"
        wb.save(f)
        df = load_with_unmerge(f)
        expected = pd.read_excel(f, header=None)
        assert df.iat[0, 0] == 10
        assert df.iat[1, 0] == 20
        # Pas pire que pandas sur la cellule de formule non cachée.
        a3, p3 = df.iat[2, 0], expected.iat[2, 0]
        assert (pd.isna(a3) and pd.isna(p3)) or a3 == p3

    def test_openpyxl_fallback_recovers_formula_text(self, tmp_path):
        # Le chemin de repli conserve le texte de la formule (jamais perdu).
        from core.compilation.merge_handler import _load_with_unmerge_openpyxl
        f = str(tmp_path / "formula2.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["A2"] = 20
        ws["A3"] = "=SUM(A1:A2)"
        ws["B3"] = "total"
        wb.save(f)
        df = _load_with_unmerge_openpyxl(f, 0)
        assert df.iat[2, 0] == "=SUM(A1:A2)"

    def test_empty_sheet_returns_empty_frame(self, tmp_path):
        f = str(tmp_path / "vide.xlsx")
        openpyxl.Workbook().save(f)
        df = load_with_unmerge(f)
        # DataFrame vide cohérent (pas une fausse cellule None).
        assert df.empty

    def test_fast_and_openpyxl_paths_agree(self, tmp_path):
        # Les deux chemins (voie rapide XML et repli openpyxl) doivent produire
        # un DataFrame identique, fusions comprises — garantie anti-divergence.
        from core.compilation.merge_handler import _load_with_unmerge_openpyxl
        f = _make_xlsx(tmp_path / "agree.xlsx", [
            ["DEP", "AUTEURS", None, None],
            [None, "Ens", "Elv", "Etu"],
            [None, 1, 2, 3],
            [None, 4, 5, 6],
        ], merges=["A1:A4", "B1:D1"])
        fast = load_with_unmerge(f)
        slow = _load_with_unmerge_openpyxl(f, 0)
        assert fast.shape == slow.shape
        assert fast.fillna("#").values.tolist() == slow.fillna("#").values.tolist()

    def test_corrupted_file_raises_clear_error(self, tmp_path):
        # Un fichier non-Excel ne doit pas crasher avec un message opaque :
        # le helper relève une ValueError explicite (le moteur l'isole ensuite).
        import pytest
        f = tmp_path / "fake.xlsx"
        f.write_text("ceci n'est pas un xlsx")
        with pytest.raises(ValueError):
            load_with_unmerge(str(f))
