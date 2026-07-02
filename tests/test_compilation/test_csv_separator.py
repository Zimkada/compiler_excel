"""
Tests de la détection du séparateur CSV (E2).

Un CSV « à la française » utilise le point-virgule (Excel FR l'exporte ainsi,
la virgule servant de séparateur décimal). Lu naïvement avec la virgule, tout
le contenu tombait dans UNE colonne (corruption silencieuse) ou faisait échouer
la détection de référence. On teste que le séparateur est correctement deviné
partout : compilation, aperçu, et détecteurs.
"""

import openpyxl
import pandas as pd

from core.compilation import ExcelCompiler, CompilationOptions, OutputFormat
from core.compilation.compilation_models import FilenameOption
from core.detection.base_detector import sniff_csv_separator


def _write_csv(path, content):
    with open(path, 'w', encoding='utf-8-sig') as f:
        f.write(content)
    return str(path)


def _opts():
    return CompilationOptions(
        use_reference_mode=True, reference_header_row=1, reference_header_lines=1,
        filename_option=FilenameOption.NONE,
    )


def _compile(tmp_path, content):
    p = _write_csv(tmp_path / 'in.csv', content)
    out = tmp_path / 'out.xlsx'
    result = ExcelCompiler(_opts()).compile_files([p], str(out), OutputFormat.XLSX)
    ws = openpyxl.load_workbook(out).active
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    return result, ws.max_column, rows


# --------------------------------------------------------------------------- #
# sniff_csv_separator (unitaire)
# --------------------------------------------------------------------------- #
class TestSniffer:
    def test_semicolon(self, tmp_path):
        p = _write_csv(tmp_path / 'f.csv', 'A;B;C\n1;2;3\n')
        assert sniff_csv_separator(p) == ';'

    def test_comma(self, tmp_path):
        p = _write_csv(tmp_path / 'f.csv', 'A,B,C\n1,2,3\n')
        assert sniff_csv_separator(p) == ','

    def test_tab(self, tmp_path):
        p = _write_csv(tmp_path / 'f.csv', 'A\tB\tC\n1\t2\t3\n')
        assert sniff_csv_separator(p) == '\t'

    def test_single_column_defaults_to_comma(self, tmp_path):
        """Une seule colonne (aucun séparateur) : repli sûr sur la virgule."""
        p = _write_csv(tmp_path / 'f.csv', 'Ligne1\nLigne2\n')
        assert sniff_csv_separator(p) == ','

    def test_empty_file_defaults_to_comma(self, tmp_path):
        p = _write_csv(tmp_path / 'f.csv', '')
        assert sniff_csv_separator(p) == ','

    def test_semicolon_wins_over_decimal_comma(self, tmp_path):
        """Piège : données à virgule décimale mais séparateur point-virgule.
        Le point-virgule (2 colonnes) doit l'emporter sur la virgule (qui
        compterait les décimales)."""
        p = _write_csv(tmp_path / 'f.csv', 'Produit;Prix\nPomme;1,50\nPoire;2,30\n')
        assert sniff_csv_separator(p) == ';'

    def test_quoted_field_with_semicolons_not_confused(self, tmp_path):
        """Cas adverse (trouvé en certification) : vrai séparateur = virgule,
        mais un champ entre guillemets contient plusieurs ';'. Le comptage
        naïf de caractères choisissait ';' à tort (le champ produisait plus de
        colonnes). csv.reader respecte les guillemets ET la consistance
        départage : ',' donne [2,2,2] (consistant), ';' donne [1,4,4]
        (erratique) -> ',' l'emporte."""
        p = _write_csv(
            tmp_path / 'f.csv',
            'Nom,Description\nProduit1,"a; b; c; d"\nProduit2,"e; f; g; h"\n')
        assert sniff_csv_separator(p) == ','

    def test_quoted_field_with_semicolons_compiles_correctly(self, tmp_path):
        """Le cas adverse ne doit plus faire ÉCHOUER la compilation."""
        result, cols, rows = _compile(
            tmp_path,
            'Nom,Description\nProduit1,"a; b; c; d"\nProduit2,"e; f; g; h"\n')
        assert result.success
        assert cols == 2
        assert rows[1] == ('Produit1', 'a; b; c; d')


# --------------------------------------------------------------------------- #
# Compilation de bout en bout
# --------------------------------------------------------------------------- #
class TestCompilation:
    def test_french_semicolon_not_corrupted(self, tmp_path):
        """Le cas prouvé : CSV français ; -> 3 colonnes, accents intacts."""
        result, cols, rows = _compile(
            tmp_path, 'Nom;Prénom;Val\nDupont;José;1\nMartin;Aïcha;2\n')
        assert result.success
        assert cols == 3
        assert rows[0] == ('Nom', 'Prénom', 'Val')
        # Accents préservés
        assert 'José' in [rows[1][1]]

    def test_comma_still_works(self, tmp_path):
        result, cols, _ = _compile(
            tmp_path, 'Name,City,Score\nAlice,Paris,10\nBob,Lyon,20\n')
        assert result.success
        assert cols == 3

    def test_decimal_comma_with_semicolon_separator(self, tmp_path):
        """Données à virgule décimale, séparées par ; : 2 colonnes, valeurs
        décimales préservées."""
        result, cols, rows = _compile(
            tmp_path, 'Produit;Prix\nPomme;1,50\nPoire;2,30\n')
        assert result.success
        assert cols == 2
        assert str(rows[1][1]) == '1,50'


# --------------------------------------------------------------------------- #
# Invariant aperçu = compilation
# --------------------------------------------------------------------------- #
class TestPreviewInvariant:
    def test_preview_sees_same_columns_as_compilation(self, tmp_path):
        p = _write_csv(tmp_path / 'in.csv', 'A;B;C\n1;2;3\n4;5;6\n')
        preview = ExcelCompiler(_opts()).preview_detection([p])[0]
        assert preview.success
        assert len(preview.detected_headers) == 3
