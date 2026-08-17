"""
Tests de non-régression du BorderDetector sur le nombre de lignes d'en-tête.

Bug corrigé : dans un tableau entièrement quadrillé, le détecteur étendait
les en-têtes sur toutes les lignes de données (header_rows = 5, 13, voire 29),
car il s'appuyait sur la seule densité de bordures. data_start_row tombait
alors après la fin du fichier => 0 ligne de données extraite.

Correctif : la fin des en-têtes est déterminée par le contenu (les lignes de
données contiennent des valeurs numériques, les en-têtes sont textuels), avec
un garde-fou pour les tableaux 100 % textuels (en-tête = 1 ligne par défaut).
"""

import openpyxl
from openpyxl.styles import Border, Side

from core.detection.border_detector import BorderDetector


_THIN = Side(style="thin")
_FULL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _make_bordered_xlsx(path, rows, border_from_row=1):
    wb = openpyxl.Workbook()
    ws = wb.active
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row, 1):
            cell = ws.cell(ri, ci, val)
            if ri >= border_from_row:
                cell.border = _FULL_BORDER
    wb.save(path)
    return str(path)


class TestBorderHeaderRows:
    def test_single_header_row_in_fully_gridded_table(self, tmp_path):
        """Tableau quadrillé, en-tête textuel + données numériques => 1 ligne."""
        p = _make_bordered_xlsx(tmp_path / "t.xlsx", [
            ["N° ordre", "Matricule", "Nom"],
            [1, 79974, "OLODO"],
            [2, 35304, "MIGAN"],
            [3, 116541, "YAROU"],
        ])
        r = BorderDetector().detect(p)
        assert r.header_start_row == 1
        assert r.header_rows == 1
        assert r.data_start_row == 2

    def test_leading_blank_rows_then_header(self, tmp_path):
        """En-tête précédé de lignes vides + titres (cas des fichiers réels)."""
        p = _make_bordered_xlsx(tmp_path / "t.xlsx", [
            [None, None, None],
            [None, None, None],
            ["Nom", "Age", "Ville"],
            ["Alice", 30, "Paris"],
            ["Bob", 25, "Lyon"],
        ], border_from_row=3)
        r = BorderDetector().detect(p)
        assert r.header_start_row == 3
        assert r.header_rows == 1
        assert r.data_start_row == 4

    def test_multiline_header(self, tmp_path):
        """En-tête réel sur 2 lignes textuelles avant des données numériques."""
        p = _make_bordered_xlsx(tmp_path / "t.xlsx", [
            ["Info", "Info"],
            ["Nom", "Age"],
            ["Alice", 30],
            ["Bob", 25],
        ])
        r = BorderDetector().detect(p)
        assert r.header_start_row == 1
        assert r.header_rows == 2

    def test_all_text_table_defaults_to_one_header_row(self, tmp_path):
        """Tableau 100 % texte : pas de signal numérique => en-tête = 1 ligne,
        les données ne doivent pas être avalées."""
        p = _make_bordered_xlsx(tmp_path / "t.xlsx", [
            ["Nom", "Ville"],
            ["Alice", "Paris"],
            ["Bob", "Lyon"],
            ["Carol", "Nice"],
        ])
        r = BorderDetector().detect(p)
        assert r.header_rows == 1
        assert r.data_start_row == 2

    def test_numeric_cell_in_header_row_is_safe(self, tmp_path):
        """Si la ligne d'en-tête contient un nombre, on garde 1 ligne (sûr)."""
        p = _make_bordered_xlsx(tmp_path / "t.xlsx", [
            ["Nom", 2024],
            ["Alice", 12],
            ["Bob", 15],
        ])
        r = BorderDetector().detect(p)
        assert r.header_rows == 1
        assert r.data_start_row == 2

    def test_single_bordered_row_does_not_crash(self, tmp_path):
        """Tableau réduit à une seule ligne bordée (header_start == data_end)."""
        p = _make_bordered_xlsx(tmp_path / "t.xlsx", [["Nom", "Age"]])
        r = BorderDetector().detect(p)
        assert r is not None
        assert r.header_rows >= 1

    def test_date_data_keeps_single_header(self, tmp_path):
        """Des données de type date ne doivent pas être prises pour des en-têtes."""
        from datetime import datetime
        p = _make_bordered_xlsx(tmp_path / "t.xlsx", [
            ["Event", "Quand"],
            ["A", datetime(2025, 3, 8)],
            ["B", datetime(2025, 4, 1)],
        ])
        r = BorderDetector().detect(p)
        assert r.header_rows == 1


class TestRowNumericCount:
    """Comportement de _row_numeric_count (utilisé pour séparer en-tête/données)."""

    def _df(self, *rows):
        import pandas as pd
        return pd.DataFrame(list(rows))

    def test_out_of_bounds_returns_zero(self):
        det = BorderDetector()
        assert det._row_numeric_count(self._df(["a", 1]), 99) == 0

    def test_numeric_string_counts(self):
        det = BorderDetector()
        assert det._row_numeric_count(self._df(["12", "x"]), 0) == 1

    def test_text_only_counts_zero(self):
        det = BorderDetector()
        assert det._row_numeric_count(self._df(["Nom", "Ville"]), 0) == 0

    def test_mixed_row(self):
        det = BorderDetector()
        # 2 numériques (1, "42"), 1 texte
        assert det._row_numeric_count(self._df([1, "Alice", "42"]), 0) == 2


class TestVerticalMergeFallback:
    """En-tête sur 2 niveaux quand les DONNÉES ne contiennent aucun nombre.

    Le critère par contenu (« une ligne de données porte des chiffres ») est
    aveugle quand la seule ligne de données vaut « NEANT », du texte libre, ou
    quand le formulaire n'est pas rempli. Le détecteur retombait alors à 1 seule
    ligne d'en-tête : les libellés n'étaient plus aplatis comme ceux du reste du
    lot (« Plants 2025 » au lieu de « Plants 2025 - Nombre total ») et
    l'aligneur créait des colonnes en double dans la compilation.

    Les fusions VERTICALES (A3:A4) décrivent la structure indépendamment du
    contenu : elles servent de repli.
    """

    def _make(self, path, data_row):
        wb = openpyxl.Workbook()
        ws = wb.active
        rows = [
            ["SUIVI DES PLANTS", None, None, None, None],
            [None, None, None, None, None],
            ["Etablissement", "Plants 2025", None, "Plants 2026", None],
            [None, "Nombre total", "Nombre survécu", "Nombre total", "Nombre survécu"],
            data_row,
        ]
        for ri, row in enumerate(rows, 1):
            for ci, val in enumerate(row, 1):
                cell = ws.cell(ri, ci, val)
                if ri >= 3:
                    cell.border = _FULL_BORDER
        # En-tête à deux niveaux : identité fusionnée verticalement,
        # catégories fusionnées horizontalement au-dessus des sous-colonnes.
        ws.merge_cells("A3:A4")
        ws.merge_cells("B3:C3")
        ws.merge_cells("D3:E3")
        wb.save(path)
        return str(path)

    def test_text_data_uses_vertical_merge(self, tmp_path):
        """Données « NEANT » : aucun nombre, mais A3:A4 dit en-tête sur 2 lignes."""
        p = self._make(tmp_path / "neant.xlsx",
                       ["CEG SAM", "NEANT", "NEANT", "NEANT", "NEANT"])
        r = BorderDetector().detect(p)
        assert r.header_start_row == 3
        assert r.header_rows == 2
        assert r.data_start_row == 5

    def test_empty_data_row_uses_vertical_merge(self, tmp_path):
        """Formulaire non rempli : même conclusion."""
        p = self._make(tmp_path / "vide.xlsx", [None, None, None, None, None])
        r = BorderDetector().detect(p)
        assert r.header_start_row == 3
        assert r.header_rows == 2

    def test_numeric_data_still_uses_content(self, tmp_path):
        """Non-régression : quand les données ont des nombres, le critère par
        contenu reste maître et donne le même résultat."""
        p = self._make(tmp_path / "chiffres.xlsx", ["CEG X", 30, 20, 35, 28])
        r = BorderDetector().detect(p)
        assert r.header_start_row == 3
        assert r.header_rows == 2
        assert r.data_start_row == 5

    def test_no_vertical_merge_keeps_single_header_row(self, tmp_path):
        """Sans fusion verticale et sans nombres, le défaut sûr (1 ligne)
        reste appliqué : on n'invente pas un en-tête sur 2 lignes."""
        p = _make_bordered_xlsx(tmp_path / "plat.xlsx", [
            ["Nom", "Ville"],
            ["Alice", "Kandi"],
        ])
        r = BorderDetector().detect(p)
        assert r.header_rows == 1
