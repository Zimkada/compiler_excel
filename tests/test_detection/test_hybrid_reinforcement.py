"""
Tests de non-régression du renforcement de la détection hybride (Phase D).

Verrouillent les causes racines diagnostiquées et corrigées :

- D1 : DensityDetector absorbait les lignes de données majoritairement
  textuelles (ex. « P0 | 0 | V0 ») dans l'en-tête -> header_rows=21,
  data_start corrompu. Corrigé par le critère de contenu numérique
  (même recette que le fix BorderDetector de juin, test_border_header_rows.py).
- D2 : l'accord de >= 2 détecteurs sur la même ligne d'en-tête était DILUÉ par
  la moyenne pondérée (0.54 et 0.73, tous deux corrects, donnaient 0.64 —
  sous le seuil 0.65 : détection correcte rejetée). Corrigé par le bonus
  d'accord (max individuel + 0.15, plafonné 0.95).
- D3 : le vote médian sur header_rows/data_end n'était pas robuste à un
  détecteur aberrant (median([21,1])=11 -> data tronquées 10/20). Corrigé :
  header_rows aberrants (>5) écartés du vote ; data_end = max (jamais de
  troncature silencieuse), 0 (= jusqu'au bout) prioritaire.
- D4 : une détection écartée sous le seuil basculait en manuel SANS un mot.
  Corrigé : warning visible dans l'aperçu et la compilation.
- D5 : les images / zones de texte (calque DrawingML) sont ignorées sans que
  l'utilisateur le sache. Corrigé : notice informative dans l'aperçu.

Fixtures 100% synthétiques (les fichiers réels ont été retirés du dépôt).
"""

import os

import openpyxl
import pandas as pd
import pytest

from core.compilation import ExcelCompiler, CompilationOptions
from core.detection.base_detector import DetectionResult, prune_phantom_columns
from core.detection.density_detector import DensityDetector
from core.detection.hybrid_detector import HybridDetector


# --------------------------------------------------------------------------- #
# Fixtures synthétiques
# --------------------------------------------------------------------------- #
def _make_table(path, header_row=5, n_data=20, textual_data=False,
                header_lines=1):
    """Tableau synthétique : lignes 1..header_row-1 vides, en-tête(s), données.

    textual_data=True produit des données SANS aucun nombre (test du garde-fou
    « tableau 100% textuel »). Sinon, données mixtes majoritairement textuelles
    (« P0 | 0 | V0 » : 2 textes / 1 nombre) — le motif qui piégeait density.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    labels = [['Nom', 'Valeur', 'Ville'], ['Prénom', 'Montant', 'Zone']]
    for li in range(header_lines):
        for c, txt in enumerate(labels[li], start=1):
            ws.cell(header_row + li, c, txt)
    first_data = header_row + header_lines
    for i in range(n_data):
        if textual_data:
            ws.cell(first_data + i, 1, f'P{chr(65 + i % 26)}')
            ws.cell(first_data + i, 2, 'texte')
            ws.cell(first_data + i, 3, f'V{chr(65 + i % 4)}')
        else:
            ws.cell(first_data + i, 1, f'P{i}')
            ws.cell(first_data + i, 2, i * 3)
            ws.cell(first_data + i, 3, f'V{i % 4}')
    wb.save(path)
    return str(path)


def _load_pruned(path):
    df = pd.read_excel(path, header=None)
    df, _ = prune_phantom_columns(df)
    return df


def _auto_options():
    return CompilationOptions(use_reference_mode=False,
                              auto_detect_structure=True)


# --------------------------------------------------------------------------- #
# D1 — DensityDetector : header_rows non gonflé
# --------------------------------------------------------------------------- #
class TestDensityHeaderRows:
    def test_not_inflated_by_text_majority_data(self, tmp_path):
        """Des données « P0 | 0 | V0 » (texte-majoritaires mais numériques)
        ne doivent PAS être absorbées dans l'en-tête (bug header_rows=21)."""
        p = _make_table(tmp_path / 't.xlsx')
        r = DensityDetector(0.5).detect(p, _load_pruned(p))
        assert r.header_start_row == 5
        assert r.header_rows == 1
        assert r.data_start_row == 6

    def test_textual_table_defaults_to_single_header(self, tmp_path):
        """Tableau 100% textuel : le critère numérique n'a pas de butoir ->
        défaut sûr : 1 ligne d'en-tête (jamais engloutir les données)."""
        p = _make_table(tmp_path / 't.xlsx', textual_data=True)
        r = DensityDetector(0.5).detect(p, _load_pruned(p))
        assert r.header_rows == 1

    def test_genuine_multiline_header_still_detected(self, tmp_path):
        """Un vrai en-tête sur 2 lignes (purement textuelles) reste détecté
        comme tel — le correctif ne doit pas casser le multi-lignes."""
        p = _make_table(tmp_path / 't.xlsx', header_lines=2)
        r = DensityDetector(0.5).detect(p, _load_pruned(p))
        assert r.header_start_row == 5
        assert r.header_rows == 2
        assert r.data_start_row == 7

    def test_dates_and_decimals_count_as_data(self, tmp_path):
        """Les dates et décimaux en données comptent bien comme numériques
        (via float()) : l'en-tête ne les absorbe pas."""
        from datetime import datetime
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A5'] = 'Client'; ws['B5'] = 'Date'; ws['C5'] = 'Montant'
        for i in range(15):
            ws.cell(6 + i, 1, f'C{i}')
            ws.cell(6 + i, 2, datetime(2025, 1, 1 + i % 28))
            ws.cell(6 + i, 3, i * 100.5)
        p = tmp_path / 'dates.xlsx'
        wb.save(p)
        r = DensityDetector(0.5).detect(str(p), _load_pruned(p))
        assert r.header_start_row == 5
        assert r.header_rows == 1

    def test_known_limit_text_only_first_data_row(self, tmp_path):
        """LIMITE CONNUE (documentée, non un bug) : si la 1re ligne de données
        est 100% textuelle (ex. « TOTAL | Ligne texte »), le critère « pas de
        nombre = en-tête » l'absorbe dans l'en-tête. Ce cas est ambigu même
        pour un humain ; le mode référence (défaut de l'app) le résout en
        fixant la ligne. On verrouille le comportement (borné, jamais
        catastrophique) plutôt que de tenter une heuristique fragile."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A3'] = 'Code'; ws['B3'] = 'Libelle'
        ws['A4'] = 'TOTAL'; ws['B4'] = 'Ligne texte'   # 1re data 100% texte
        for i in range(10):
            ws.cell(5 + i, 1, i)
            ws.cell(5 + i, 2, f'item{i}')
        p = tmp_path / 'mix.xlsx'
        wb.save(p)
        r = DensityDetector(0.5).detect(str(p), _load_pruned(p))
        # header_rows peut valoir 2 (absorbe la ligne texte) mais reste borné
        # et n'engloutit jamais les vraies données numériques (à partir de L5).
        assert r.header_start_row == 3
        assert r.header_rows <= 5
        assert r.data_start_row <= 5  # les données numériques ne sont pas mangées

    def test_header_rows_hard_cap(self, tmp_path):
        """Même en présence de nombreuses lignes textuelles consécutives,
        l'en-tête est plafonné à 5 lignes."""
        wb = openpyxl.Workbook()
        ws = wb.active
        # 8 lignes purement textuelles puis des données numériques
        for i in range(8):
            ws.cell(1 + i, 1, f'Libellé{i}')
            ws.cell(1 + i, 2, f'Sous{i}')
        for i in range(10):
            ws.cell(9 + i, 1, f'P{i}')
            ws.cell(9 + i, 2, i)
        p = tmp_path / 'cap.xlsx'
        wb.save(p)
        r = DensityDetector(0.5).detect(str(p), _load_pruned(p))
        assert r.header_rows <= 5


# --------------------------------------------------------------------------- #
# D2 — Bonus d'accord inter-détecteurs (scénario prouvé n°1)
# --------------------------------------------------------------------------- #
class TestAgreementBonus:
    def test_simple_table_passes_threshold(self, tmp_path):
        """Scénario prouvé : tableau simple, en-tête ligne 5. Density (0.54)
        et pattern (0.73) s'accordaient sur la ligne 5 mais la moyenne (0.64)
        passait sous le seuil 0.65 -> détection correcte rejetée. L'accord
        doit maintenant produire une confiance >= seuil ET la bonne ligne."""
        p = _make_table(tmp_path / 't.xlsx')
        preview = ExcelCompiler(_auto_options()).preview_detection([p])[0]
        assert preview.success
        assert preview.header_start_row == 5
        assert preview.confidence >= 0.65
        assert preview.detection_method == 'hybrid'
        # Toutes les données présentes (pas de troncature, cf. D3)
        assert preview.data_row_count == 20

    def test_agreement_recorded_in_debug(self, tmp_path):
        p = _make_table(tmp_path / 't.xlsx')
        r = HybridDetector().detect(p, _load_pruned(p))
        assert r.debug_info.get('selection_reason') == 'detector_agreement'
        assert r.confidence <= 0.95  # plafond du bonus


# --------------------------------------------------------------------------- #
# D3 — Vote robuste (unitaires sur _create_combined_result)
# --------------------------------------------------------------------------- #
def _res(method, conf, start, rows=1, end=0):
    return DetectionResult(
        file_path='x.xlsx', detection_method=method, confidence=conf,
        header_start_row=start, header_rows=rows, data_start_row=start + rows,
        data_end_row=end,
    )


class TestRobustVoting:
    WEIGHTS = {'border': 0.40, 'density': 0.30, 'pattern': 0.30}

    def test_aberrant_header_rows_excluded_from_vote(self):
        """header_rows=[21, 1] : la médiane donnait 11 (troncature 10/20).
        L'aberrant (>5) doit être écarté -> vote = 1."""
        results = [_res('density', 0.54, 5, rows=21, end=26),
                   _res('pattern', 0.73, 5, rows=1, end=25)]
        combined = HybridDetector()._create_combined_result(results, self.WEIGHTS)
        assert combined.header_rows == 1
        assert combined.data_start_row == 6

    def test_all_aberrant_defaults_to_one(self):
        results = [_res('density', 0.6, 5, rows=21),
                   _res('pattern', 0.6, 5, rows=30)]
        combined = HybridDetector()._create_combined_result(results, self.WEIGHTS)
        assert combined.header_rows == 1

    def test_data_end_takes_max_never_truncates(self):
        """data_end=[15, 25] : la médiane coupait à 20 ; on prend le MAX (25).
        Garder trop est bénin (filtres aval), tronquer est destructeur."""
        results = [_res('density', 0.6, 5, end=15),
                   _res('pattern', 0.6, 5, end=25)]
        combined = HybridDetector()._create_combined_result(results, self.WEIGHTS)
        assert combined.data_end_row == 25

    def test_data_end_zero_means_until_end_and_wins(self):
        """0 = « jusqu'au bout » : si un détecteur d'accord le dit, aucune
        borne ne doit être imposée."""
        results = [_res('density', 0.6, 5, end=0),
                   _res('pattern', 0.6, 5, end=25)]
        combined = HybridDetector()._create_combined_result(results, self.WEIGHTS)
        assert combined.data_end_row == 0

    def test_disagreement_keeps_weighted_average(self):
        """Sans accord sur header_start, le comportement historique
        (moyenne pondérée + médiane) est conservé."""
        results = [_res('density', 0.54, 3),
                   _res('pattern', 0.73, 7)]
        combined = HybridDetector()._create_combined_result(results, self.WEIGHTS)
        assert combined.debug_info['selection_reason'] == 'weighted_vote'
        # Moyenne pondérée à poids égaux (0.3/0.3) = (0.54+0.73)/2
        assert abs(combined.confidence - 0.635) < 0.01


# --------------------------------------------------------------------------- #
# D4 — Transparence du rejet sous seuil
# --------------------------------------------------------------------------- #
class TestRejectionTransparency:
    def test_preview_warns_when_detection_rejected(self, tmp_path):
        """Avec un seuil inatteignable (0.99), la détection est écartée : le
        basculement en manuel doit être SIGNALÉ, pas silencieux."""
        p = _make_table(tmp_path / 't.xlsx')
        opts = _auto_options()
        opts.detection_confidence_threshold = 0.99
        preview = ExcelCompiler(opts).preview_detection([p])[0]
        assert preview.detection_method == 'manual'
        assert preview.warning is not None
        assert 'écartée' in preview.warning

    def test_compilation_warns_when_detection_rejected(self, tmp_path):
        from core.compilation import OutputFormat
        p = _make_table(tmp_path / 't.xlsx')
        opts = _auto_options()
        opts.detection_confidence_threshold = 0.99
        out = tmp_path / 'out.xlsx'
        result = ExcelCompiler(opts).compile_files([p], str(out), OutputFormat.XLSX)
        assert any('écartée' in w for w in result.warnings)

    def test_no_rejection_warning_when_detection_used(self, tmp_path):
        """Quand la détection est acceptée, aucun faux warning de rejet."""
        p = _make_table(tmp_path / 't.xlsx')
        preview = ExcelCompiler(_auto_options()).preview_detection([p])[0]
        assert preview.detection_method == 'hybrid'
        assert not (preview.warning and 'écartée' in preview.warning)


# --------------------------------------------------------------------------- #
# D5 — Notice « formes ignorées » (nécessite Pillow pour créer l'image)
# --------------------------------------------------------------------------- #
class TestIgnoredShapesNotice:
    def test_preview_notices_drawings(self, tmp_path):
        PILImage = pytest.importorskip('PIL.Image')
        from openpyxl.drawing.image import Image as XLImage

        logo = tmp_path / 'logo.png'
        PILImage.new('RGB', (60, 30), (0, 100, 0)).save(logo)

        p = tmp_path / 'avec_logo.xlsx'
        _make_table(p)
        wb = openpyxl.load_workbook(p)
        img = XLImage(str(logo))
        img.anchor = 'A1'
        wb.active.add_image(img)
        wb.save(p)

        preview = ExcelCompiler(_auto_options()).preview_detection([str(p)])[0]
        assert preview.success
        assert preview.warning and 'zones de texte' in preview.warning
        # Scénario prouvé n°2 : plus aucune troncature avec logo (20/20)
        assert preview.data_row_count == 20
        assert preview.header_start_row == 5

    def test_no_notice_without_drawings(self, tmp_path):
        p = _make_table(tmp_path / 'sans.xlsx')
        preview = ExcelCompiler(_auto_options()).preview_detection([p])[0]
        assert not (preview.warning and 'zones de texte' in preview.warning)
