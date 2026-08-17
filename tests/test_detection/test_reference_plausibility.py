"""
Avertissement quand la ligne d'en-tête indiquée n'est pas un en-tête.

Cas vécu : l'utilisateur laisse « ligne 1 » alors que le tableau commence
ligne 3. La ligne 1 ne porte que le TITRE du document (« SUIVI DES PLANTS DES
CAMPAGNES… », fusionné sur toute la largeur). Tous les fichiers du lot portant
ce même titre, la similarité affiche 100 % et rien ne signale l'erreur : la
compilation « réussit » en transformant les en-têtes réels en lignes de
données.

Le contrôle est volontairement conservateur — il ne bloque rien, il avertit.
"""

import openpyxl
import pytest

from core.detection.reference_detector import ReferenceDetector


def _make(path, rows, merges=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row, 1):
            if val is not None:
                ws.cell(ri, ci, val)
    for m in merges:
        ws.merge_cells(m)
    wb.save(path)
    return str(path)


FORMULAIRE = [
    ["SUIVI DES PLANTS DES CAMPAGNES 2025 ET 2026", None, None, None, None],
    [None, None, None, None, None],
    ["Etablissement", "Plants 2025", None, "Plants 2026", None],
    [None, "Nombre total", "Nombre survécu", "Nombre total", "Nombre survécu"],
    ["CEG ARBONGA", 189, 189, 72, 68],
]


class TestReferencePlausibility:
    def test_title_row_triggers_warning(self, tmp_path):
        p = _make(tmp_path / "form.xlsx", FORMULAIRE, merges=("A1:E1",))
        d = ReferenceDetector(reference_file=p, reference_header_row=1,
                              reference_header_lines=1)
        assert d.reference_warning is not None
        assert "ligne 1" in d.reference_warning
        assert "titre du document" in d.reference_warning

    def test_real_header_row_is_silent(self, tmp_path):
        p = _make(tmp_path / "form.xlsx", FORMULAIRE, merges=("A1:E1",))
        d = ReferenceDetector(reference_file=p, reference_header_row=3,
                              reference_header_lines=2)
        assert d.reference_warning is None

    def test_partial_header_row_is_silent(self, tmp_path):
        """Ligne 3 seule : 3 libellés réels sur 5 colonnes — plausible."""
        p = _make(tmp_path / "form.xlsx", FORMULAIRE, merges=("A1:E1",))
        d = ReferenceDetector(reference_file=p, reference_header_row=3,
                              reference_header_lines=1)
        assert d.reference_warning is None

    def test_narrow_table_never_warns(self, tmp_path):
        """Un tableau de 2-3 colonnes peut légitimement n'avoir qu'un ou deux
        libellés : on ne se prononce qu'à partir de 4 colonnes."""
        p = _make(tmp_path / "etroit.xlsx", [
            ["Nom", None],
            ["Alice", 12],
        ])
        d = ReferenceDetector(reference_file=p, reference_header_row=1,
                              reference_header_lines=1)
        assert d.reference_warning is None

    def test_two_labels_on_wide_table_warns(self, tmp_path):
        """Deux libellés pour 6 colonnes reste suspect."""
        p = _make(tmp_path / "large.xlsx", [
            ["Rapport annuel", "2026", None, None, None, None],
            ["Etab", "A", "B", "C", "D", "E"],
            ["CEG X", 1, 2, 3, 4, 5],
        ])
        d = ReferenceDetector(reference_file=p, reference_header_row=1,
                              reference_header_lines=1)
        assert d.reference_warning is not None

    def test_warning_is_advisory_not_blocking(self, tmp_path):
        """Le détecteur reste pleinement fonctionnel malgré l'avertissement."""
        p = _make(tmp_path / "form.xlsx", FORMULAIRE, merges=("A1:E1",))
        d = ReferenceDetector(reference_file=p, reference_header_row=1,
                              reference_header_lines=1)
        assert d.reference_warning is not None
        result = d.detect(p)
        assert result.header_start_row == 1
        assert result.confidence > 0
